"""Documents for Karen: read Word/PDF/PowerPoint/text, create Word/PDF, organise folders."""

import re
import shutil
from collections import Counter
from pathlib import Path

from agent import resolve

CHUNK = 6000  # characters per read; keeps a request inside Groq's free token limit
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css", ".java", ".c", ".cpp",
                   ".h", ".log", ".xml", ".yml", ".yaml", ".ini", ".dart", ".kt", ".sql", ".tex", ".rtf"}


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------
def _docx_text(p: Path) -> tuple[str, dict]:
    import docx

    d = docx.Document(p)
    parts = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in block.iter() if t.tag.endswith("}t"))
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iter():
                        if cell.tag.endswith("}tc"):
                            cells.append("".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t")).strip())
                    parts.append(" | ".join(cells))
    return "\n".join(parts), {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}


def _pdf_text(p: Path) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        pages.append(f"--- page {i} ---\n{text}")
    text = "\n".join(pages)
    info = {"pages": len(reader.pages)}
    if len(re.sub(r"--- page \d+ ---|\s", "", text)) < 20 * max(1, len(reader.pages)):
        info["note"] = "this PDF is mostly images (scanned); little text could be extracted"
    return text, info


def _pptx_text(p: Path) -> tuple[str, dict]:
    from pptx import Presentation

    prs = Presentation(str(p))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        slides.append(f"--- slide {i} ---\n" + "\n".join(t for t in texts if t) + (f"\n[notes] {notes}" if notes else ""))
    return "\n".join(slides), {"slides": len(prs.slides)}


def extract_text(p: Path) -> tuple[str, dict]:
    ext = p.suffix.lower()
    if ext == ".docx":
        return _docx_text(p)
    if ext == ".pdf":
        return _pdf_text(p)
    if ext == ".pptx":
        return _pptx_text(p)
    if ext in (".doc", ".ppt"):
        raise ValueError(f"old {ext} files aren't supported; save it as {ext}x first")
    raw = p.read_bytes()
    if b"\x00" in raw[:4000]:
        raise ValueError(f"{p.name} isn't a document Karen can read")
    return raw.decode("utf-8", "replace"), {}


def read_document(path: str, start: int = 0) -> dict:
    p = resolve(path)
    if not p.is_file():
        return {"error": f"no file at {p}"}
    import vision

    if vision.is_image(p.name):
        return vision.describe_image(str(p), "Describe this image and read out all the text in it.")
    try:
        text, info = extract_text(p)
    except ValueError as e:
        return {"error": str(e)}
    start = max(0, int(start or 0))
    chunk = text[start:start + CHUNK]
    result = {"path": str(p), "total_characters": len(text), **info, "start": start, "content": chunk}
    if start + CHUNK < len(text):
        result["more"] = True
        result["next_start"] = start + CHUNK
        result["hint"] = "For a summary of a long document, read the next parts with next_start before answering."
    return result


# ---------------------------------------------------------------------------
# Creating
# ---------------------------------------------------------------------------
def _blocks(content: str):
    """Very small markdown: '# ' / '## ' / '### ' headings, '- ' or '* ' bullets, '1. ' numbers, paragraphs."""
    for line in content.replace("\r\n", "\n").split("\n"):
        s = line.rstrip()
        if not s.strip():
            yield ("blank", "")
        elif m := re.match(r"^(#{1,3})\s+(.*)", s):
            yield (f"h{len(m.group(1))}", m.group(2))
        elif m := re.match(r"^\s*[-*•]\s+(.*)", s):
            yield ("bullet", m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", s):
            yield ("number", m.group(1))
        else:
            yield ("para", s)


def _inline_bold(text: str):
    """Split '**bold**' runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            yield part, i % 2 == 1


def _make_docx(p: Path, title: str, content: str) -> None:
    import docx

    d = docx.Document()
    if title:
        d.add_heading(title, level=0)
    for kind, text in _blocks(content):
        if kind == "blank":
            continue
        if kind.startswith("h"):
            d.add_heading(text, level=int(kind[1]))
            continue
        style = {"bullet": "List Bullet", "number": "List Number"}.get(kind)
        para = d.add_paragraph(style=style)
        for part, bold in _inline_bold(text):
            para.add_run(part).bold = bold
    d.save(p)


def _make_pdf(p: Path, title: str, content: str) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    if title:
        story += [Paragraph(escape(title), styles["Title"]), Spacer(1, 0.3 * cm)]
    items: list = []
    list_kind = None

    def flush():
        nonlocal items, list_kind
        if items:
            story.append(ListFlowable(items, bulletType="bullet" if list_kind == "bullet" else "1", leftIndent=14))
            items, list_kind = [], None

    def markup(text):
        return "".join(f"<b>{escape(t)}</b>" if b else escape(t) for t, b in _inline_bold(text))

    for kind, text in _blocks(content):
        if kind in ("bullet", "number"):
            if list_kind not in (None, kind):
                flush()
            list_kind = kind
            items.append(ListItem(Paragraph(markup(text), styles["BodyText"])))
            continue
        flush()
        if kind == "blank":
            story.append(Spacer(1, 0.2 * cm))
        elif kind.startswith("h"):
            story.append(Paragraph(escape(text), styles[{"h1": "Heading1", "h2": "Heading2", "h3": "Heading3"}[kind]]))
        else:
            story.append(Paragraph(markup(text), styles["BodyText"]))
    flush()
    SimpleDocTemplate(str(p), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, title=title).build(story)


def create_document(path: str, content: str, title: str = "", confirmed: bool = False) -> dict:
    p = resolve(path)
    ext = p.suffix.lower()
    if ext not in (".docx", ".pdf"):
        return {"error": "use a .docx or .pdf file name (for plain text use write_file)"}
    if p.exists() and not confirmed:
        return {"needs_confirmation": True, "will_do": f"replace the existing document {p}"}
    p.parent.mkdir(parents=True, exist_ok=True)
    (_make_docx if ext == ".docx" else _make_pdf)(p, title, content)
    return {"created": str(p), "size_kb": round(p.stat().st_size / 1024, 1)}


# ---------------------------------------------------------------------------
# Organising
# ---------------------------------------------------------------------------
CATEGORIES = {
    "Documents": {".pdf", ".doc", ".docx", ".txt", ".rtf", ".odt", ".md", ".ppt", ".pptx", ".xls", ".xlsx", ".csv"},
    "Images": {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg", ".heic", ".ico"},
    "Videos": {".mp4", ".mkv", ".mov", ".avi", ".webm", ".wmv"},
    "Audio": {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"},
    "Archives": {".zip", ".rar", ".7z", ".tar", ".gz"},
    "Installers": {".exe", ".msi", ".apk", ".msix"},
    "Code": {".py", ".js", ".ts", ".html", ".css", ".java", ".c", ".cpp", ".ipynb", ".json", ".dart", ".kt", ".sql"},
}


def _category(p: Path) -> str:
    ext = p.suffix.lower()
    return next((name for name, exts in CATEGORIES.items() if ext in exts), "Others")


def organize_folder(path: str = "downloads", confirmed: bool = False) -> dict:
    folder = resolve(path)
    if not folder.is_dir():
        return {"error": f"no folder at {folder}"}
    files = [f for f in folder.iterdir() if f.is_file() and not f.name.startswith((".", "~$"))
             and f.name.lower() != "desktop.ini"]
    if not files:
        return {"folder": str(folder), "note": "no loose files to organise"}
    plan = Counter(_category(f) for f in files)
    if not confirmed:
        return {"needs_confirmation": True, "will_do": f"move {len(files)} files in {folder} into subfolders: "
                + ", ".join(f"{n} {c}" for c, n in plan.most_common())}
    moved = 0
    for f in files:
        target_dir = folder / _category(f)
        target_dir.mkdir(exist_ok=True)
        target = target_dir / f.name
        n = 2
        while target.exists():
            target = target_dir / f"{f.stem} ({n}){f.suffix}"
            n += 1
        try:
            shutil.move(str(f), str(target))
            moved += 1
        except OSError:
            pass  # file in use: leave it where it is
    return {"organised": str(folder), "moved_files": moved, "into": dict(plan)}
